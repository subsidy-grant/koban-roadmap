# Hallmark · macrostructure: Manifesto(deck) · tone: brutal-coach · anchor hue: vermilion(中) on kinari paper
# theme: custom (vibe: "生成りの地・墨の文字・中の朱" · paper #F5F1E6 · accent #C23B1E · Dela Gothic One + Zen Kaku Gothic New)
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
import os

HERE = os.path.dirname(os.path.abspath(__file__))
MEDIA = os.path.join(HERE, "media")

# ---- tokens ----
PAPER = RGBColor(0xF5, 0xF1, 0xE6)     # 生成り（明るい地）
PAPER2 = RGBColor(0xE3, 0xEC, 0xE5)    # 淡いフェルト色の面
LINE = RGBColor(0xC9, 0xC2, 0xAE)      # 罫
INK = RGBColor(0x1B, 0x2A, 0x24)       # 墨（深緑寄り）
MUTED = RGBColor(0x4E, 0x5F, 0x57)     # 補助文字
ACCENT = RGBColor(0xC2, 0x3B, 0x1E)    # 中の朱
ON_ACCENT = RGBColor(0xFF, 0xFF, 0xFF)  # 朱の上の文字
TILE = RGBColor(0xFF, 0xFF, 0xFF)      # 牌の面
TILE_INK = RGBColor(0x1C, 0x1F, 0x1B)  # 牌面の文字
TILE_EDGE = RGBColor(0xB8, 0xAE, 0x94)
DISPLAY = "Dela Gothic One"
BODY = "Zen Kaku Gothic New"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = 13.333, 7.5
BLANK = prs.slide_layouts[6]
TOTAL = 18


def set_font(run, face, size, color, bold=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", face)


def bg(slide, color=PAPER):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, spacing=1.45, gap=6):
    """paras: list of paragraphs; each paragraph = list of (text, face, size, color[, bold])."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(gap)
        if isinstance(runs, tuple):
            runs = [runs]
        for r in runs:
            run = p.add_run()
            run.text = r[0]
            set_font(run, r[1], r[2], r[3], r[4] if len(r) > 4 else False)
    return tb


def B(t, size=17, color=INK, bold=False):
    return (t, BODY, size, color, bold)


def D(t, size=40, color=INK):
    return (t, DISPLAY, size, color)


def folio(slide, n, label):
    text(slide, W - 3.6, H - 0.55, 3.1, 0.3,
         [[B(label + "　", 11, MUTED), B(f"{n:02d} / {TOTAL}", 11, MUTED, True)]], align=PP_ALIGN.RIGHT)


def bullets(items, size=17, color=INK, mark="■", mark_color=ACCENT):
    out = []
    for it in items:
        if isinstance(it, str):
            it = [B(it, size, color)]
        out.append([(mark + "  ", BODY, size * 0.6, mark_color)] + it)
    return out


def tiles(slide, x, y, chars, tw=0.62, th=0.84, reds=()):
    """牌の並びを描く。chars の ' ' は間隔。"""
    cx = x
    for i, c in enumerate(chars):
        if c == " ":
            cx += tw * 0.45
            continue
        t = rect(slide, cx, y, tw, th, TILE, TILE_EDGE, MSO_SHAPE.ROUNDED_RECTANGLE)
        t.adjustments[0] = 0.12
        rect(slide, cx, y + th - 0.1, tw, 0.1, RGBColor(0xD9, 0xA4, 0x41))  # 牌の背（山吹）
        text(slide, cx, y + 0.02, tw, th - 0.12, [[(c, DISPLAY, 26, ACCENT if i in reds else TILE_INK)]],
             anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, spacing=1.0, gap=0)
        cx += tw + 0.06
    return cx


def rack(slide, x, y, img, w=6.0):
    """手牌画像を象牙の台に載せる（元画像は白牌＋灰文字なので明るい台が要る）。"""
    h = w * 120 / 620
    rect(slide, x - 0.18, y - 0.14, w + 0.36, h + 0.28, TILE, LINE)
    slide.shapes.add_picture(os.path.join(MEDIA, img), Inches(x), Inches(y), Inches(w), Inches(h))
    return y + h + 0.14


def verdict(slide, x, y, label, strong):
    fill, ink = (ACCENT, ON_ACCENT) if strong else (INK, PAPER)
    r = rect(slide, x, y, 0.2 + 0.52 * len(label), 0.72, fill)
    text(slide, x, y, 0.2 + 0.52 * len(label), 0.72, [[(label, DISPLAY, 28, ink)]],
         anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, spacing=1.0, gap=0)
    return x + 0.2 + 0.52 * len(label)


# =========================================================== 01 表紙
s = prs.slides.add_slide(BLANK); bg(s)
rect(s, 0, 0, 0.28, H, ACCENT)
text(s, 1.0, 0.8, 8, 0.5, [[B("博太の", 22, MUTED, True)]])
text(s, 1.0, 1.35, 11.5, 2.4, [[D("麻雀", 118), D("再生工場", 118)]], spacing=1.0)
text(s, 1.0, 3.45, 3, 2.2, [[D("2", 150, ACCENT)]], spacing=1.0)
# 副露編 の判子
st = rect(s, 2.75, 4.1, 3.3, 1.25, PAPER, ACCENT)
st.line.width = Pt(4)
text(s, 2.75, 4.1, 3.3, 1.25, [[D("副露編", 52, ACCENT)]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, spacing=1.0, gap=0)
text(s, 1.0, 6.45, 6, 0.4, [[B("有料講座", 14, MUTED, True)]])
tiles(s, 8.6, 5.9, "東南西北中", reds=(4,))

# =========================================================== 02 始めに
s = prs.slides.add_slide(BLANK); bg(s)
text(s, 0.9, 0.7, 5, 0.9, [[D("始めに", 40, ACCENT)]])
text(s, 0.9, 1.9, 6.3, 3.6, [[D("俺より上手くない", 44)], [D("と思うなら、", 44)], [D("面前でいけ。", 44, ACCENT)]], spacing=1.2, gap=0)
rect(s, 7.7, 1.9, 0.04, 4.2, LINE)
text(s, 8.1, 1.9, 4.5, 4.6, [
    [B("鳴き、副露は判断が非常に難しく、知識と経験と技術が必要。")],
    [B("だからこそ、使いこなせる人は和了に近づくし、下手な人は放銃率が高くなるだけ。")],
    [B("俺個人は直近ではむしろ鳴かないように意識している。")],
    [B("俺より麻雀が上手くはないと感じているのであれば、面前主体でいくことをオススメする。", 17, INK, True)],
], spacing=1.6, gap=10)
folio(s, 2, "副露編")

# =========================================================== 03 幻想を捨てろ
s = prs.slides.add_slide(BLANK); bg(s)
text(s, 0.9, 0.9, 6.0, 3.2, [[D("幻想を", 64)], [D("捨てろ、", 64)], [D("自我を", 64)], [D("捨てろ。", 64, ACCENT)]], spacing=1.05, gap=0)
text(s, 7.4, 0.95, 5.2, 5.6, [
    [B("鳴いて聴牌に近づける、一発消し、かわし手等々。それらをやれるほど、あなたはそんなに上手くない。")],
    [B("牌理、牌効率、牌の安全度等々を理解していないヤツが副露しても、大怪我をするだけ。")],
    [B("副露してあがれると思うな。", 24, ACCENT, True)],
    [B("自分の手で聴牌して、ツモれる形を作ることだけ考えろ！", 17, INK, True)],
    [B("自分の中のセオリーは捨てろ！", 17, INK, True)],
], spacing=1.6, gap=12)
folio(s, 3, "副露編")

# =========================================================== 04 鳴くな！
s = prs.slides.add_slide(BLANK); bg(s)
text(s, 0.9, 0.6, 12, 2.4, [[D("鳴くな！", 130, ACCENT)]], spacing=1.0)
text(s, 0.9, 2.95, 11.5, 0.8, [[B("攻撃、守備ともに知識・技術・戦術がなく、いつも負け組になっているあなたは、以下の2つを実践していくべき。", 18, MUTED)]])
for i, (num, t) in enumerate([("①", "鳴いてはいけない"), ("②", "面前でリーチをかけていく")]):
    x = 0.9 + i * 6.0
    rect(s, x, 3.95, 5.6, 1.9, PAPER2)
    rect(s, x, 3.95, 0.1, 1.9, ACCENT)
    text(s, x + 0.45, 4.1, 5.0, 1.6, [[D(num + " ", 28, ACCENT), D(t, 24)]], anchor=MSO_ANCHOR.MIDDLE, spacing=1.2)
text(s, 0.9, 6.25, 6, 0.6, [[D("なぜか　→", 26, MUTED)]])
folio(s, 4, "副露編")

# =========================================================== 05/06 デメリット・メリット
def pro_con(n, kind, sub, rows, concl, note=None):
    s = prs.slides.add_slide(BLANK); bg(s)
    text(s, 0.9, 0.6, 12, 1.2, [[D("鳴くな！", 40, ACCENT), D("　" + kind, 40)]])
    text(s, 0.9, 1.45, 12, 0.5, [[B(sub, 15, MUTED)]])
    y = 2.15
    for head, arrow, extra in rows:
        rect(s, 0.9, y, 11.5, 0.03, LINE)
        text(s, 0.9, y + 0.2, 5.6, 1.2, [[B(head, 21, INK, True)]] + ([[B(extra, 14, MUTED)]] if extra else []), spacing=1.4, gap=2)
        text(s, 6.8, y + 0.2, 5.6, 1.2, [[("→  ", BODY, 21, ACCENT, True), B(arrow, 19, ACCENT, True)]], spacing=1.4)
        y += 1.45
    rect(s, 0.9, 5.2, 11.5, 1.45, PAPER2)
    rect(s, 0.9, 5.2, 0.1, 1.45, ACCENT)
    text(s, 1.25, 5.2, 1.6, 1.45, [[D("結論", 26, ACCENT)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 2.7, 5.2, 9.5, 1.45, concl, anchor=MSO_ANCHOR.MIDDLE, spacing=1.45, gap=2)
    if note:
        text(s, 0.9, 6.8, 9, 0.35, [[B(note, 12, MUTED)]])
    folio(s, n, "副露編")

pro_con(5, "デメリット編", "鳴くと失うもの", [
    ("安全牌を持ちながらの進行が難しくなる", "守備力低下、放銃率の上昇", "Mリーガーの手牌進行をみると、安全牌を抱えて進めているのがよくわかる"),
    ("打点が上昇しづらくなる", "リーチの役と裏ドラを放棄することになる", None),
], [[B("守備力低下に見合う打点があれば鳴いてもいいが、その判断ができないうちはデメリットが大きいので", 17, INK)],
    [B("鳴くな！（あなたのことです！）", 20, ACCENT, True)]],
    "ちなみに、なぜ安全牌を持ちながら進行していくと良いのかは、また別料金。")

pro_con(6, "メリット編", "面前リーチで得るもの", [
    ("他家が押し返しづらくなる", "降ろす or 回らせる → リーチ者だけが和了の抽選を受ける「一人旅」ができる", None),
    ("打点上昇が見込める", "リーチの一翻と裏ドラが捲れるだけで他家への脅威となる（ノミ手は別）", "三麻だと大抵はドラ、赤、花が絡むので満貫以上になりやすい"),
], [[B("相手が攻めづらくなり、かつ自分の打点を上げることができる", 17, INK)],
    [B("最大最強の武器がリーチ。", 22, ACCENT, True)]])

# =========================================================== 07 これから
s = prs.slides.add_slide(BLANK); bg(s)
text(s, 0.9, 0.6, 5, 0.9, [[D("これから", 40, ACCENT)]])
text(s, 0.9, 1.7, 6.8, 5.0, [
    [B("今のあなたの副露状況は、", 19, MUTED)],
    [B("稼ぐ手段が未熟なのに、支出の機会だけ増えている状態。", 26, INK, True)],
    [B("稼ぐスキルを磨くべきだが、そのためには鳴いてはいけない。", 18)],
    [B("当然、副露を使えばあがれたであろう場面に沢山遭遇する。負けが増える期間も多少あるかもしれない（現に増えてるからいいか）。", 16, MUTED)],
], spacing=1.55, gap=12)
rect(s, 8.2, 1.7, 4.3, 3.5, PAPER2)
text(s, 8.55, 1.95, 3.7, 3.1, [
    [B("それよりも、リーチしてツモるための手組、そのための牌理・牌効率・場況を考え、技術を向上させることに重きを置いて上達することが", 16)],
    [B("勝ち組への近道。", 20, ACCENT, True)],
], spacing=1.6, gap=8)
text(s, 8.2, 5.5, 4.5, 1.2, [[D("だから", 26, MUTED)], [D("鳴くな！", 48, ACCENT)]], spacing=1.0, gap=0)
folio(s, 7, "副露編")

# =========================================================== 08 それでも
s = prs.slides.add_slide(BLANK); bg(s)
text(s, 0.9, 0.6, 12, 0.9, [[D("それでも", 40, ACCENT)]])
text(s, 0.9, 1.45, 12, 0.5, [[B("鳴きたい時は基準や条件を持つ（三麻と四麻でも当然変わる）", 16, MUTED)]])
cols = [
    ("鳴く条件", ACCENT, ON_ACCENT, ["跳満・満貫以上が見込めて、聴牌形も両面以上の良形で先制できるとき",
                               "手牌の中に他家の現物が3つ以上、あるいは字牌トイツがあるとき"]),
    ("鳴かない条件", INK, PAPER, ["鳴いてもシャンテン数が上がらない、聴牌まで遠い",
                                  "鳴いても愚形（ペンチャン、カンチャン、シャンポン、単騎）が残る"]),
]
for i, (h, fill, hink, items) in enumerate(cols):
    x = 0.9 + i * 5.95
    rect(s, x, 2.15, 5.55, 0.8, fill)
    text(s, x + 0.3, 2.15, 5, 0.8, [[D(h, 26, hink)]], anchor=MSO_ANCHOR.MIDDLE)
    rect(s, x, 2.95, 5.55, 2.55, PAPER2)
    text(s, x + 0.35, 3.2, 4.9, 2.2, bullets(items, 18, mark_color=fill), spacing=1.5, gap=12)
text(s, 0.9, 5.8, 11.6, 1.0, [[B("攻撃・守備ともに知識・技術・戦術がない負け組のあなたは、鳴く時にはこれぐらいの高い基準を持つように。", 15, INK, True)],
                              [B("鳴くまでの条件をフローチャートにして作成してみるのもいいかも。", 15, MUTED)]], spacing=1.5, gap=2)
folio(s, 8, "副露編")

# =========================================================== 09 番外編1 鳴け！
s = prs.slides.add_slide(BLANK); bg(s)
text(s, 0.9, 0.55, 12, 1.1, [[B("番外編 1　", 18, MUTED, True), D("鳴け！", 44, ACCENT)]], anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.9, 1.7, 12, 0.45, [[B("三麻　東一局　西家　ドラ⑤　6巡目", 16, MUTED, True)]])
tiles(s, 0.9, 2.2, "南南中中中 334789 ⑨⑨", tw=0.66, th=0.9, reds=(2, 3, 4))
text(s, 0.9, 3.35, 11.6, 0.8, [[B("リーチしてもリーチ・中のみだが、南、⑨はポンして", 18), B("先制両面聴牌", 18, ACCENT, True), B("が取れるイーシャンテン。", 18)]])
text(s, 0.9, 4.15, 11.6, 2.8, bullets([
    [B("ドラが絡まないので打点上昇する確率は低い。"), B("親を流すことを優先", 17, INK, True)],
    [B("自分にドラがないので"), B("他家の打点が高くなっている", 17, INK, True), B("と予想し、かわしていく")],
    "これでリーチを打っても、他家のドラ爆とのめくりあいは期待値が低い",
    "他家の先制リーチが入っていた場合、現物切って聴牌できなければオリ",
]), spacing=1.45, gap=8)
folio(s, 9, "番外編")

# =========================================================== 10 番外編1 ホンイツは？
s = prs.slides.add_slide(BLANK); bg(s)
text(s, 0.9, 0.55, 12, 1.1, [[B("番外編 1　", 18, MUTED, True), D("鳴け！", 44, ACCENT)]], anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.9, 1.8, 12, 1.0, [[D("南ポンして⑨切りのホンイツは？", 34)]])
text(s, 0.9, 2.85, 7.2, 4.0, bullets([
    [B("最低でも（都合良く入ってきたとして）もう一巡かかるため", 16), B("却下", 16, ACCENT, True)],
    [B("時間をかければかけるほど、他家の手もよくなっていくことを忘れるな", 16, INK, True)],
    [B("4のくっつきではあるが待ちが良くなりづらい（嬉しい変化は自分で2枚使っている3と5だけ）", 16)],
    [B("面前、他家のスピードがなさそう、かつ⑨が他2人に対しての安全牌（になりそう）ならホンイツを採用してもよし。", 16, MUTED),
     B("そんなの分からないだろうから採用しないが吉。", 16, INK, True)],
], 16), spacing=1.5, gap=10)
rect(s, 8.6, 2.85, 3.9, 3.8, ACCENT)
text(s, 8.95, 3.05, 3.3, 3.4, [
    [B("スピードと打点のバランス、待ち、親流しのメリットと比較した時に、どう考えても釣り合わない。", 15, ON_ACCENT)],
    [B("少しでも頭をよぎったら", 17, ON_ACCENT, True)],
    [D("死刑", 64, ON_ACCENT)],
], spacing=1.45, gap=6)
folio(s, 10, "番外編")

# =========================================================== 11 番外編2 回し打ち
s = prs.slides.add_slide(BLANK); bg(s)
text(s, 0.9, 0.55, 12, 1.1, [[B("番外編 2　", 18, MUTED, True), D("回し打ち（絶対）やめろ", 44, ACCENT)]], anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.9, 1.75, 11.5, 0.8, [[B("下手なヤツが回し打ちすると、途中から欲かいて危険牌バンバン切り出すから", 18), B("絶対にやめろ。", 18, ACCENT, True)]])
for i, (h, sub) in enumerate([("ゼンツ", "全ツッパ"), ("ベタオリ", "現物抜き打って")]):
    x = 0.9 + i * 3.35
    rect(s, x, 2.7, 3.05, 1.7, INK if i == 0 else PAPER2)
    text(s, x, 2.8, 3.05, 1.5, [[B(sub, 14, PAPER if i == 0 else MUTED, True)], [D(h, 40, PAPER if i == 0 else INK)]],
         anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, spacing=1.2, gap=2)
text(s, 0.9, 4.55, 6.4, 0.5, [[B("どちらかに振り切れ。", 20, INK, True)]])
text(s, 0.9, 5.15, 6.3, 1.8, [[B("その時にも点棒状況とシャンテン数と打点、聴牌形を考慮しろ。統計的にも出てるから、それは自分で覚えろ。", 15, MUTED)]], spacing=1.5)
rect(s, 7.75, 2.7, 0.04, 4.0, LINE)
text(s, 8.15, 2.7, 4.4, 4.1, [
    [B("いつもそれで痛い目見てることを自覚しろ。", 17)],
    [B("低打点の愚形待ちゼンツして、いじられて煽られてキレるな。", 17)],
    [D("チンパンプレイしてる", 26, ACCENT)],
    [D("自分を恥じろ！", 26, ACCENT)],
], spacing=1.5, gap=10)
folio(s, 11, "番外編")

# =========================================================== 12 番外編2 ちなみに
s = prs.slides.add_slide(BLANK); bg(s)
text(s, 0.9, 0.55, 12, 1.1, [[B("番外編 2　", 18, MUTED, True), D("ちなみに", 44, ACCENT)]], anchor=MSO_ANCHOR.MIDDLE)
paras = [
    ("Iさんの場合", "Iさんは相手の待ちを読むのが上手いゆえに、回し打ちをして攻めてくる傾向があるが、それはIさんだからできること。"),
    ("対策", "けれどもリーチをかけていればそれなりに遠回りさせられるし、おそらく本人にも自信があるからこそ聴牌や一向聴では勝負牌を切り出してくる。自分が良形聴牌していれば、Iさんからはあがりやすいと言える。"),
    ("結論", "相手の待ちは（マイティの場合特に）結論本人しかわからないので、真似して中途半端な知識で精度の低い回し打ちをするより、目に見えている確実なものだけで進行するほうが成績は安定する。"),
]
for i, (h, t) in enumerate(paras):
    x = 0.9 + i * 3.95
    rect(s, x, 1.95, 3.65, 0.06, ACCENT if i == 2 else LINE)
    text(s, x, 2.2, 3.65, 0.6, [[D(h, 24, ACCENT if i == 2 else INK)]])
    text(s, x, 2.95, 3.65, 3.9, [[B(t, 16, INK if i == 2 else MUTED, i == 2)]], spacing=1.7)
folio(s, 12, "番外編")

# =========================================================== 13/14 おまけ 問題
def question_slide(n, qs, lead):
    s = prs.slides.add_slide(BLANK); bg(s)
    text(s, 0.9, 0.5, 8.5, 1.1, [[B("おまけ　", 18, MUTED, True), D("ポンする？しない？", 44)]], anchor=MSO_ANCHOR.MIDDLE)
    rect(s, 10.3, 0.62, 2.2, 0.85, ACCENT)
    text(s, 10.3, 0.62, 2.2, 0.85, [[B("1問", 14, ON_ACCENT, True), D("30秒", 26, ON_ACCENT)]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, spacing=1.0, gap=0)
    text(s, 0.9, 1.5, 9, 0.4, [[B(lead, 14, MUTED)]])
    y = 2.0
    for label, meta, call, img in qs:
        text(s, 0.9, y, 1.9, 1.0, [[D(label, 36, ACCENT)]], spacing=1.0)
        text(s, 2.7, y + 0.02, 10, 0.5, [[B(meta, 16, MUTED, True)]])
        text(s, 2.7, y + 0.42, 10, 0.6, [[B(call, 22, INK, True)]])
        rack(s, 2.88, y + 1.1, img, 5.8)
        y += 2.6
    folio(s, n, "おまけ")

question_slide(13, [
    ("①", "三麻　東一局　西家　3巡目　35000点持ち", "東家から 東（2枚目）", "image3.png"),
    ("②", "三麻　東三局　西家　5巡目　40000点持ち", "東家から 7索", "image2.png"),
], "全4問")
question_slide(14, [
    ("③-1", "三麻　東三局　南家　5巡目　50000点持ち", "東家から 東（1枚目）", "image1.png"),
    ("③-2", "三麻　東三局　南家　5巡目　50000点持ち", "東家から 東（1枚目）", "image4.png"),
], "③-1 と ③-2 は手牌の違いに注目")

# =========================================================== 15-18 解答
def answer_slide(n, label, meta, img, badge, strong, answer_rest, pts, closing):
    s = prs.slides.add_slide(BLANK); bg(s)
    text(s, 0.9, 0.5, 5, 0.9, [[B("解答　", 18, MUTED, True), D(label, 40, ACCENT)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 0.9, 1.35, 7.5, 0.4, [[B(meta, 14, MUTED, True)]])
    rack(s, 1.08, 1.95, img, 5.9)
    bx = verdict(s, 7.65, 1.95, badge, strong)
    text(s, 7.65, 2.85, 4.9, 0.9, [[B(answer_rest, 15, MUTED)]], spacing=1.45)
    text(s, 0.9, 3.55, 6.6, 3.4, bullets(pts, 14.5), spacing=1.5, gap=6)
    rect(s, 7.65, 4.0, 4.85, 2.75, PAPER2)
    rect(s, 7.65, 4.0, 0.08, 2.75, ACCENT)
    text(s, 8.0, 4.15, 4.3, 2.45, [[B(closing, 17, INK, True)]], anchor=MSO_ANCHOR.MIDDLE, spacing=1.55)
    folio(s, n, "おまけ・解答")

answer_slide(15, "①", "三麻　東一局　西家　3巡目　35000点持ち　東家から東（2枚目）", "image3.png",
    "鉄板でしない", False, "親だったら3翻の手で受け入れ枚数が増えるからする。これでポンすると答えた人は麻雀のセンス０（俺の主観だけどね）",
    ["2枚目といえどカンチャンが2つ残っている、2翻の二向聴。鳴いた後も雀頭候補が無く、他家のリーチに押し返しづらい牌姿のため、安全牌候補の一萬九萬を残しつつ、面前聴牌のリーチ自摸ドラ1を目指す。",
     "ピンズも④で三面張、⑤⑧でもリャンカンが残る。索子もまだ横に伸びていける形。"],
    "現状二向聴、鳴いても二向聴。打点が上がらず待ちも良くならない最悪の鳴きであることを理解できないとやばいからね？")

answer_slide(16, "②", "三麻　東三局　西家　5巡目　40000点持ち　東家から7索", "image2.png",
    "絶対にしない", False, "俺は絶対にしない（オーラスアガリトップならする）",
    ["鳴いて4000点の聴牌は安全にとれるが、カンチャンが残り、手牌が中張牌だけになるため他家のリーチに押し返しづらい。",
     "索子は四連形で横に伸びやすく、ピンズは①③④引くと両面、⑥はダイレクト聴牌（ダマで変化待ち）、ドラ⑨を引いて②⑤⑧のくっつき聴牌などへの良形変化が多くあるため、面前で進める。",
     "ピンズの一通など打点も見込め、打点があれば他家のリーチに押し返しやすい。最悪7索が親の現物になる。"],
    "今は、鳴いて打点を下げて愚形待ちにしないように。")

answer_slide(17, "③-1", "三麻　東三局　南家　5巡目　50000点持ち　東家から東（1枚目）", "image1.png",
    "しない推奨", False, "微妙だが、（微妙だからこそ）しない推奨。2枚目ならポン",
    ["赤5索のくっつきやホンイツで跳満以上も見える手牌だが、最終的に嵌4索、嵌6索や、出づらい⑤とドラの白とのシャンポン待ちになる可能性もある。打点があるからそれでもいいけど、今は良形聴牌を目指せ。",
     "赤5索とホンイツの両天秤で進めると、他家のリーチが来てからのホンイツが行きづらく、ドラを切らないとなると中張牌で溢れるリスクがある。"],
    "まだ4ブロックしかないターツ不足と、赤5索が確実に使えるか分からない点が判断を難しくさせている。点棒はあるため、判断が微妙な時は無理しない。")

answer_slide(18, "③-2", "三麻　東三局　南家　5巡目　50000点持ち　東家から東（1枚目）", "image4.png",
    "鉄板でする", True, "③-1 との違いは、赤5索が確実に使えること",
    ["③-1と同じく4ブロックだが、異なるのは赤5索が確実に使え、5ブロック目は四連形のピンズで作れそうな良い形になっていること。",
     "これだけの好材料で跳満も見えるなら、点棒状況関係なく攻める一手。",
     "危惧しているシャンポン待ちへの変化は②⑤の6枚、両面以上は①③④⑥の14枚。②ならまだ出やすいので、ピンズは5種17枚が嬉しい変化となる。⑦もターツ不足解消の牌なので、良形変化するまでは持っておく。"],
    "打点があり、愚形でないなら全ツッパ！")

out = os.path.join(HERE, "博太の麻雀再生工場2_副露編.pptx")
prs.save(out)
print(out)
